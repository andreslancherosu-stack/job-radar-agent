"""Deterministic tools for the discovery-and-report run.

Division of labour, enforced here rather than merely described:

  Python decides   what was retrieved, whether a quote is real, whether a
                   requisition is already in the ledger, what the workbook
                   says, which URLs got opened.
  The model decides whether a posting suits the user, how to rank what it
                   found, what is still unresolved, and what to do next.

There is deliberately NO fit score, no weighting table and no ranking
function in this file. A tool that computed suitability and asked the model
to restate it would make the model decorative, which is exactly the failure
mode this project is graded against.

Run-scoped state lives in RUN. It is reset per run so a tool can refer to a
posting by a short key ("gartner:112849") instead of the model having to
quote whole postings back -- that is a context-cost decision, not cosmetic.
"""
import datetime
import json
import re
import webbrowser
from pathlib import Path

import discovery
import tools as core_tools

REPORTS_DIR = Path('private/reports')

MAX_SELECTIONS = 4

# Visible slice of a search result. See search_openings for why this is a
# recurring cost rather than a one-off.
MAX_CANDIDATES_SHOWN = 45


class Run:
    """Everything one radar run accumulates. Reset per run, never global state
    that leaks between runs."""

    def __init__(self):
        self.candidates = {}        # "source:job_id" -> compact row
        self.details = {}           # "source:job_id" -> full detail incl. full text
        self.full_text = {}         # "source:job_id" -> text used for the evidence gate
        self.coverage = []          # what was ACTUALLY reached, per source
        self.selections = []        # model's picks, validated
        self.searches = 0
        self.report_path = None
        self.opened_urls = []
        self.started = datetime.datetime.now()

    def key(self, row):
        return '%s:%s' % (row.get('source'), row.get('job_id'))

    def record_coverage(self, entries):
        """Union this search's coverage into the run's. A source counts as
        reached if ANY search reached it; scanned/kept totals add up; an error
        is kept only while that source has never succeeded."""
        by_source = {c['source']: c for c in self.coverage}
        for e in entries:
            prev = by_source.get(e['source'])
            if prev is None:
                by_source[e['source']] = dict(e, searches=1)
                continue
            prev['searches'] = prev.get('searches', 1) + 1
            prev['returned_by_source'] += e['returned_by_source']
            prev['kept_after_filter'] += e['kept_after_filter']
            if e['reached']:
                prev['reached'] = True
                prev['error'] = None if prev.get('reached') else prev.get('error')
                if e.get('total_reported_by_source') is not None:
                    prev['total_reported_by_source'] = e['total_reported_by_source']
            elif not prev['reached']:
                prev['error'] = e.get('error') or prev.get('error')
        self.coverage = sorted(by_source.values(), key=lambda c: c['source'])


RUN = Run()


def reset_run():
    global RUN
    RUN = Run()
    return RUN


def _today():
    return datetime.date.today().isoformat()


# --------------------------------------------------------------------------
# tools
# --------------------------------------------------------------------------

def list_sources():
    """What this run is ABLE to search, and what it deliberately is not."""
    by_key, raw = discovery.load_sources()
    return {
        'sources': [{'key': s['key'], 'company': s['company'], 'ats': s['ats'],
                     'on_priority_list': s.get('on_priority_list', False),
                     'verified_on': s.get('verified_on'), 'note': s.get('note')}
                    for s in by_key.values()],
        'coverage_limits': raw.get('known_coverage_limits'),
        'probed_and_rejected': [p['target'] for p in raw.get('probed_and_rejected', [])],
        'rule': ('This is the complete list of employers reachable programmatically in this '
                 'run. It is NOT the user\'s 51-company priority list. Any conclusion drawn '
                 'must be scoped to these employers only.'),
    }


def search_openings(source_keys=None, keywords=None, locations=None, limit_per_source=60,
                    search_text=None):
    """Live search across official employer APIs. Returns COMPACT rows.

    The keyword/location filter is a mechanical string match so the model
    receives tens of candidates rather than hundreds. It encodes no opinion
    about suitability: it cannot see the profile.
    """
    by_key, _ = discovery.load_sources()
    source_keys = source_keys or list(by_key)
    rows, coverage = discovery.search_all(
        source_keys, keywords=keywords, locations=locations, search_text=search_text,
        limit_per_source=limit_per_source, sources=by_key)
    RUN.searches += 1
    # ACCUMULATE, never replace. Measured bug (2026-10-11): the model ran a
    # second, narrower search near the end of the run; that overwrote
    # RUN.coverage, and the report's "actual search coverage" sheet then
    # listed a single source -- under-reporting four employers that really had
    # been searched. Coverage is a claim about the whole run, so it has to be
    # the union of every search the run performed.
    RUN.record_coverage(coverage)
    # Every candidate row returned here is re-sent to the model on EVERY
    # subsequent step, so this list is a recurring cost, not a one-off.
    # Measured: broadening coverage from 17 to 156 candidates pushed one
    # 9-step run from 168k to 274k input tokens. The full set stays available
    # to later tools via RUN.candidates; only the visible slice is capped, and
    # the cap announces itself so a truncated list is never mistaken for the
    # whole result.
    shown = []
    for r in rows:
        k = RUN.key(r)
        RUN.candidates[k] = r
        if len(shown) < MAX_CANDIDATES_SHOWN:
            row = {'key': k, 'company': r['company'], 'title': r['title'],
                   'location': r['location']}
            if r.get('deadline'):
                row['deadline'] = r['deadline']
            shown.append(row)
    reached = [c for c in coverage if c['reached']]
    omitted = len(rows) - len(shown)
    return {
        'candidates': shown,
        'candidate_count': len(rows),
        'candidates_shown': len(shown),
        'candidates_omitted': omitted,
        'omission_note': (
            'Showing the first %d of %d matches to keep this observation small. The other %d '
            'were found and are NOT lost -- narrow `keywords` or `locations`, or set a more '
            'specific `search_text`, to surface different ones.' % (len(shown), len(rows), omitted)
        ) if omitted else None,
        'coverage': coverage,
        'sources_reached': len(reached),
        'sources_failed': [c['source'] for c in coverage if not c['reached']],
        'scanned_total': sum(c['returned_by_source'] for c in coverage),
        'note': ('Titles and locations only -- no requirement text yet. Call '
                 'get_posting_details with the keys you care about (batched, several at '
                 'once) to read requirements and any sponsorship language. '
                 'A row appearing here means the employer\'s API listed it right now; it '
                 'is not by itself proof the vacancy is open to new applicants.'),
    }


def get_posting_details(keys, full_text=False):
    """Batch-fetch requirement text for several postings in ONE tool call."""
    if not keys:
        return {'error': 'keys must be a non-empty list of candidate keys from search_openings'}
    if isinstance(keys, str):
        keys = [keys]
    if len(keys) > 8:
        return {'error': 'at most 8 keys per call (got %d) -- batch, but keep observations '
                         'small enough to stay useful' % len(keys)}
    unknown = [k for k in keys if k not in RUN.candidates]
    if unknown:
        return {'error': 'unknown candidate key(s) %s -- they must come from a '
                         'search_openings result in THIS run' % unknown}
    rows = [RUN.candidates[k] for k in keys]
    details = discovery.fetch_details(rows, full_text=full_text)
    out = []
    for k, d in zip(keys, details):
        if d.get('error'):
            out.append({'key': k, 'error': d['error']})
            continue
        RUN.details[k] = d
        if full_text:
            RUN.full_text[k] = d.get('full_text', '')
        else:
            # Keep a copy of the complete text for the evidence gate even
            # though the model only saw an excerpt -- so a legitimate quote
            # from the excerpt validates, and a fabricated one still fails.
            full = discovery.fetch_detail(RUN.candidates[k], full_text=True)
            RUN.full_text[k] = full.get('full_text', '')
        entry = {'key': k, 'company': d.get('company'), 'title': d.get('title'),
                 'location': d.get('location'), 'url': d.get('url'),
                 'posted': d.get('updated'), 'deadline_field': d.get('deadline_field'),
                 'still_listed': d.get('still_listed'), 'can_apply': d.get('can_apply')}
        if full_text:
            entry['full_text'] = d.get('full_text')
        else:
            entry.update({kk: d.get(kk) for kk in (
                'sponsorship_sentences', 'deadline_sentences', 'requirement_lines',
                'requirement_lines_omitted', 'full_text_chars', 'shown_chars',
                'omission_note')})
        out.append(entry)
    return {'postings': out, 'count': len(out)}


def select_opportunity(key, priority, why_suitable, evidence_quote,
                       unresolved_questions=None, next_action=None):
    """Record ONE opportunity the model judges suitable. Python validates; it
    does not re-judge.

    Three gates, each a real failure mode this project has already hit:
      * the quote must appear VERBATIM in the retrieved posting
      * the requisition must not already be in the canonical ledger
      * at most MAX_SELECTIONS, so a run cannot pad the report
    """
    if key not in RUN.details:
        return {'error': 'call get_posting_details for %r before selecting it -- a selection '
                         'must rest on retrieved text, not on a title' % key}
    if len(RUN.selections) >= MAX_SELECTIONS:
        return {'error': 'already selected the maximum of %d opportunities' % MAX_SELECTIONS}
    if any(s['key'] == key for s in RUN.selections):
        return {'error': '%r is already selected' % key}

    d = RUN.details[key]
    quote = (evidence_quote or '').strip()
    if len(quote) < 12:
        return {'error': 'evidence_quote must be a substantive verbatim span from the posting'}
    haystack = re.sub(r'\s+', ' ', RUN.full_text.get(key, ''))
    if re.sub(r'\s+', ' ', quote) not in haystack:
        return {'error': 'Evidence not found verbatim in the retrieved posting for %s. Quote '
                         'exactly what the employer wrote, or do not claim it.' % key}

    hist, hist_err = core_tools.execute_tool('check_application_history', {
        'company': d.get('company'), 'job_id': str(d.get('job_id') or ''),
        'source_url': d.get('url')})
    if not hist_err and hist.get('already_applied'):
        return {'error': 'DUPLICATE: the ledger already records an application to %s (%s). Do '
                         'not select it.' % (key, hist.get('application_record'))}

    sel = {'key': key, 'company': d.get('company'), 'title': d.get('title'),
           'job_id': d.get('job_id'), 'location': d.get('location'), 'url': d.get('url'),
           'priority': priority, 'why_suitable': why_suitable, 'evidence_quote': quote,
           'unresolved_questions': unresolved_questions or [],
           'next_action': next_action,
           'posted': d.get('updated'), 'deadline_field': d.get('deadline_field'),
           'still_listed': d.get('still_listed'), 'can_apply': d.get('can_apply')}
    RUN.selections.append(sel)
    return {'recorded': True, 'key': key, 'selections_so_far': len(RUN.selections),
            'remaining_slots': MAX_SELECTIONS - len(RUN.selections),
            'duplicate_check': 'passed -- not in the canonical ledger'}


def export_report(summary=None, partial=False, partial_reason=None):
    """Dated Excel report. Deterministic; the model supplies no numbers."""
    import openpyxl
    from openpyxl.styles import Font, Alignment

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    stamp = RUN.started.strftime('%Y-%m-%d_%H%M%S')
    name = 'job_radar_report_%s%s.xlsx' % (stamp, '_PARTIAL' if partial else '')
    path = REPORTS_DIR / name

    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical='top')

    ws = wb.active
    ws.title = 'Selected'
    status = ('PARTIAL -- %s' % (partial_reason or 'run ended before completing')) if partial \
        else 'COMPLETE'
    ws.append(['Job Radar - selected opportunities'])
    ws['A1'].font = bold
    ws.append(['Generated', RUN.started.strftime('%Y-%m-%d %H:%M:%S')])
    ws.append(['Run status', status])
    ws.append(['Selections', len(RUN.selections)])
    ws.append(['Evaluation basis',
               'Live retrieval from official employer APIs during this run. Listing by an '
               'employer API is not confirmation that a vacancy accepts new applicants.'])
    if summary:
        ws.append(['Model summary', summary])
    ws.append([])
    head = ['Priority', 'Company', 'Title', 'Req ID', 'Location', 'Why suitable (model)',
            'Evidence (verbatim from posting)', 'Unresolved questions', 'Next action',
            'Posted', 'Deadline field', 'Still listed', 'Official application URL']
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = bold
    for s in sorted(RUN.selections, key=lambda x: str(x.get('priority'))):
        ws.append([s.get('priority'), s.get('company'), s.get('title'), s.get('job_id'),
                   s.get('location'), s.get('why_suitable'), s.get('evidence_quote'),
                   '; '.join(s.get('unresolved_questions') or []) or 'none stated',
                   s.get('next_action'), str(s.get('posted') or ''),
                   str(s.get('deadline_field') or ''), str(s.get('still_listed')),
                   s.get('url')])
    for col, width in zip('ABCDEFGHIJKLM', (9, 14, 34, 11, 20, 52, 58, 40, 34, 16, 14, 12, 56)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=8):
        for c in row:
            c.alignment = wrap

    cs = wb.create_sheet('Search coverage')
    cs.append(['What this run ACTUALLY searched'])
    cs['A1'].font = bold
    cs.append(['Searches run', RUN.searches])
    cs.append([])
    cs.append(['Source', 'Company', 'ATS', 'Reached', 'Error', 'Returned', 'Total reported',
               'Kept after keyword/location filter'])
    for c in cs[cs.max_row]:
        c.font = bold
    for c in RUN.coverage:
        cs.append([c['source'], c['company'], c['ats'], 'yes' if c['reached'] else 'NO',
                   c.get('error') or '', c['returned_by_source'],
                   c.get('total_reported_by_source'), c['kept_after_filter']])
    cs.append([])
    _, raw = discovery.load_sources()
    cs.append(['Documented coverage limits'])
    cs[cs.max_row][0].font = bold
    for lim in raw.get('known_coverage_limits', []):
        cs.append([lim])
    for col, width in zip('ABCDEFGH', (14, 16, 12, 9, 26, 10, 14, 30)):
        cs.column_dimensions[col].width = width

    wb.save(path)
    RUN.report_path = str(path)
    return {'report_path': str(path), 'selections': len(RUN.selections),
            'status': status, 'partial': bool(partial),
            'sources_reached': sum(1 for c in RUN.coverage if c['reached']),
            'sources_total': len(RUN.coverage),
            'verify': 'File written and closed. Re-read it to confirm contents.'}


def open_application_urls(keys=None):
    """Open the official application URLs for the user. Opens pages only --
    it cannot fill or submit anything."""
    sel = RUN.selections if not keys else [s for s in RUN.selections if s['key'] in keys]
    if not sel:
        return {'error': 'nothing selected to open'}
    opened, failed = [], []
    for s in sel:
        url = s.get('url')
        if not url:
            failed.append({'key': s['key'], 'error': 'no url'})
            continue
        try:
            webbrowser.open_new_tab(url)
            opened.append({'key': s['key'], 'company': s['company'], 'url': url})
        except Exception as e:
            failed.append({'key': s['key'], 'error': str(e)[:80]})
    RUN.opened_urls = opened
    return {'opened': opened, 'opened_count': len(opened), 'failed': failed,
            'note': 'Browser tabs opened to the employer application pages. No form was '
                    'filled and nothing was submitted.'}


# --------------------------------------------------------------------------
# schemas
# --------------------------------------------------------------------------

TOOL_DEFINITIONS = [
    {'name': 'list_sources',
     'description': 'List the official employer job APIs this run can search, plus documented '
                    'coverage limits. Call this first if you need to know what coverage means.',
     'input_schema': {'type': 'object', 'properties': {}}},
    {'name': 'search_openings',
     'description': 'Search live openings across official employer APIs. Returns compact rows '
                    '(title, location, key) -- NOT requirement text. Keyword and location '
                    'filters are mechanical string matches used to keep the result small; they '
                    'make no suitability judgment.',
     'input_schema': {'type': 'object', 'properties': {
         'source_keys': {'type': 'array', 'items': {'type': 'string'},
                         'description': 'Source keys from list_sources. Omit to search all.'},
         'keywords': {'type': 'array', 'items': {'type': 'string'},
                      'description': 'Match any of these in the job title.'},
         'locations': {'type': 'array', 'items': {'type': 'string'},
                       'description': 'Match any of these in the location string.'},
         'limit_per_source': {'type': 'integer',
                              'description': 'Max postings to pull per source (default 60).'},
         'search_text': {'type': 'string',
                         'description': 'ONE short term sent to the employer\'s own search (e.g. '
                                        '"analyst"). Workday narrows server-side on this, and a '
                                        'single term works far better than several: measured, one '
                                        'term kept 32 candidates where an empty term kept 7 and '
                                        'nine joined keywords kept almost none. Defaults to your '
                                        'first keyword. Your full keyword list is still applied '
                                        'locally as an OR filter.'}}}},
    {'name': 'get_posting_details',
     'description': 'Batch-fetch requirement text, sponsorship language and posting dates for '
                    'up to 8 candidates in ONE call. Pass several keys at once. Observations '
                    'are excerpted and say what was omitted; set full_text=true for one '
                    'posting when you need everything.',
     'input_schema': {'type': 'object', 'properties': {
         'keys': {'type': 'array', 'items': {'type': 'string'},
                  'description': 'Candidate keys from search_openings, e.g. "gartner:112849".'},
         'full_text': {'type': 'boolean',
                       'description': 'Return the entire posting instead of an excerpt.'}},
         'required': ['keys']}},
    {'name': 'select_opportunity',
     'description': 'Record one opportunity you judge suitable. Requires a verbatim quote from '
                    'the retrieved posting as evidence; the quote is checked against the real '
                    'text and the requisition is checked against the application ledger. '
                    'Maximum 4 selections.',
     'input_schema': {'type': 'object', 'properties': {
         'key': {'type': 'string'},
         'priority': {'type': 'integer', 'description': '1 = apply first.'},
         'why_suitable': {'type': 'string',
                          'description': 'Your reasoning, referencing the user\'s verified facts.'},
         'evidence_quote': {'type': 'string',
                            'description': 'Verbatim span from the posting supporting the match.'},
         'unresolved_questions': {'type': 'array', 'items': {'type': 'string'},
                                  'description': 'What is still genuinely unknown.'},
         'next_action': {'type': 'string'}},
         'required': ['key', 'priority', 'why_suitable', 'evidence_quote']}},
    {'name': 'export_report',
     'description': 'Write the dated Excel report with selections, evidence, unresolved '
                    'questions, official URLs and actual search coverage. The coverage sheet is '
                    'filled in automatically from what was actually reached -- you do not supply '
                    'those numbers.',
     'input_schema': {'type': 'object', 'properties': {
         'summary': {'type': 'string', 'description': 'Short overall summary for the report.'},
         'partial': {'type': 'boolean',
                     'description': 'TRUE ONLY IF THIS RUN WAS CUT SHORT -- you ran out of steps '
                                    'or time, or a source failed, so work you intended to do did '
                                    'not happen. It marks the filename _PARTIAL and stamps the '
                                    'file INCOMPLETE. It is NOT the flag for "coverage is narrow" '
                                    'or "questions remain": this tool only ever searches five '
                                    'employers and open questions are normal, so neither of those '
                                    'makes a run partial. Put those in summary and in each '
                                    'selection\'s unresolved_questions instead. If you finished '
                                    'what you set out to do, pass false.'},
         'partial_reason': {'type': 'string',
                            'description': 'Required when partial is true: what stopped the run.'}}}},
    {'name': 'open_application_urls',
     'description': 'Open the official application pages for the selected opportunities in the '
                    'browser. Opens pages only; cannot fill or submit anything.',
     'input_schema': {'type': 'object', 'properties': {
         'keys': {'type': 'array', 'items': {'type': 'string'},
                  'description': 'Omit to open every selection.'}}}},
]

TOOL_REGISTRY = {
    'list_sources': list_sources,
    'search_openings': search_openings,
    'get_posting_details': get_posting_details,
    'select_opportunity': select_opportunity,
    'export_report': export_report,
    'open_application_urls': open_application_urls,
}
