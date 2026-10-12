"""Tests for the terminal menu and the ledger-driven report rebuild.

Everything runs against a TEMPORARY state file in a directory whose name
contains spaces. No real opportunity is ever marked submitted, no paid search
runs, and the live browser session is never touched.
"""
import datetime
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import ledger_report
import menu
import radar


def posting(company, role, job_id, url, kind='employer'):
    return {'company': company, 'role': role, 'job_id': job_id, 'source_url': url,
            'source_kind': kind, 'text': 'Requirements: a real snapshot of posting text. ' * 6}


class Fixture(unittest.TestCase):
    """A ledger in a directory with SPACES in the name -- the real project
    lives under 'FALL 26 BUAN 6V99 Agentic AI', so a path-quoting bug would
    be invisible in a tidy temp dir."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='job radar test ')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / 'my project folder'
        (self.root / 'private' / 'reports').mkdir(parents=True)
        self.state_path = self.root / 'private' / 'state.json'
        self.reports = self.root / 'private' / 'reports'

        self.state = radar.empty()
        self.day = '2026-10-12'
        radar.ingest(self.state, posting('Acme', 'Data Analyst', '111',
                                         'https://acme.example/jobs/111'), '2026-10-01')
        radar.ingest(self.state, posting('Borealis', 'BI Analyst', '222',
                                         'https://borealis.example/jobs/222'), '2026-10-02')
        radar.write(self.state_path, self.state)

        for target, value in (('STATE_PATH', self.state_path),):
            p = mock.patch.object(menu, target, value)
            p.start()
            self.addCleanup(p.stop)
        for target, value in (('STATE_PATH', self.state_path), ('REPORTS_DIR', self.reports)):
            p = mock.patch.object(ledger_report, target, value)
            p.start()
            self.addCleanup(p.stop)

    def key_of(self, company, job_id):
        return radar.key({'company': company, 'job_id': job_id})

    def answers(self, *values):
        """Feed scripted console input."""
        return mock.patch.object(menu, 'ask', side_effect=list(values))


class SelectingAJob(Fixture):
    def test_lists_jobs_and_selects_by_number_not_by_key(self):
        state = radar.read(self.state_path)
        with mock.patch.object(menu, 'ask', side_effect=['1']):
            key = menu.choose_job(state)
        self.assertEqual(key, self.key_of('Acme', '111'))

    def test_zero_cancels(self):
        state = radar.read(self.state_path)
        with mock.patch.object(menu, 'ask', side_effect=['0']):
            self.assertIsNone(menu.choose_job(state))

    def test_out_of_range_is_refused_not_guessed(self):
        state = radar.read(self.state_path)
        with mock.patch.object(menu, 'ask', side_effect=['99']):
            self.assertIsNone(menu.choose_job(state))

    def test_pasted_url_resolves_to_a_saved_posting(self):
        state = radar.read(self.state_path)
        with mock.patch.object(menu, 'ask', side_effect=['u', 'https://borealis.example/jobs/222/']):
            key = menu.choose_job(state)
        self.assertEqual(key, self.key_of('Borealis', '222'))

    def test_unknown_url_is_not_registered_without_consent(self):
        state = radar.read(self.state_path)
        with mock.patch.object(menu, 'ask', side_effect=['u', 'https://new.example/jobs/999']), \
             mock.patch.object(menu, 'ask_yes_no', return_value=False):
            self.assertIsNone(menu.choose_job(state))
        self.assertEqual(len(radar.read(self.state_path)['jobs']), 2, 'nothing may be added')

    def test_unknown_url_without_retrievable_text_saves_nothing(self):
        """The ledger stores a real snapshot. Rather than invent one, the menu
        refuses -- an empty posting would be fabricated metadata."""
        state = radar.read(self.state_path)
        with mock.patch.object(menu, 'ask',
                               side_effect=['u', 'https://new.example/jobs/999',
                                            'Newco', 'Analyst', '999']), \
             mock.patch.object(menu, 'ask_yes_no', return_value=True), \
             mock.patch.object(menu, 'fetch_posting_text', return_value=None), \
             mock.patch('builtins.input', side_effect=['END']):
            self.assertIsNone(menu.choose_job(state))
        self.assertEqual(len(radar.read(self.state_path)['jobs']), 2)

    def test_unknown_url_with_real_text_uses_radars_own_ingest(self):
        state = radar.read(self.state_path)
        with mock.patch.object(menu, 'ask',
                               side_effect=['u', 'https://new.example/jobs/999',
                                            'Newco', 'Analyst', '999']), \
             mock.patch.object(menu, 'ask_yes_no', return_value=True), \
             mock.patch.object(menu, 'fetch_posting_text',
                               return_value='Real retrieved posting text. ' * 20):
            key = menu.choose_job(state)
        self.assertEqual(key, self.key_of('Newco', '999'))
        saved = radar.read(self.state_path)['jobs'][key]
        self.assertEqual(saved['posting']['source_kind'], 'user_copy')
        self.assertNotIn(key, radar.read(self.state_path)['applications'],
                         'registering a posting must NOT record an application')


class RecordingAnApplication(Fixture):
    def test_cancelling_at_the_confirmation_question_records_nothing(self):
        with mock.patch.object(menu, 'ask', side_effect=['1']), \
             mock.patch.object(menu, 'ask_yes_no', return_value=False):
            menu.record_application()
        self.assertEqual(radar.read(self.state_path)['applications'], {})

    def test_unfinished_verification_is_not_a_submission(self):
        """Clicking Submit with an email-code step still pending is exactly the
        case the question is worded to exclude."""
        with mock.patch.object(menu, 'ask', side_effect=['1']), \
             mock.patch.object(menu, 'ask_yes_no', side_effect=[False]):
            menu.record_application()
        self.assertEqual(radar.read(self.state_path)['applications'], {})

    def test_successful_recording_goes_through_radar_confirm(self):
        key = self.key_of('Acme', '111')
        with mock.patch.object(menu, 'ask', side_effect=['1', '2026-10-12']), \
             mock.patch.object(menu, 'ask_yes_no', side_effect=[True, True]), \
             mock.patch.object(radar, 'confirm', wraps=radar.confirm) as spy, \
             mock.patch.object(ledger_report, 'build', return_value={
                 'report_path': 'x', 'csv_path': 'y', 'pending': 1, 'submitted': 1}):
            menu.record_application()
        self.assertTrue(spy.called, 'must use radar.confirm, not a parallel writer')
        saved = radar.read(self.state_path)['applications']
        self.assertIn(key, saved)
        self.assertEqual(saved[key]['date_applied'], '2026-10-12')

    def test_the_date_defaults_visibly_to_today_in_chicago(self):
        seen = {}

        def fake_ask(prompt, default=None):
            if 'submission date' in prompt.lower():
                seen['prompt'] = prompt
                seen['default'] = default
                return default
            return '1'

        with mock.patch.object(menu, 'ask', side_effect=fake_ask), \
             mock.patch.object(menu, 'ask_yes_no', side_effect=[True, True]), \
             mock.patch.object(ledger_report, 'build', return_value={
                 'report_path': 'x', 'csv_path': 'y', 'pending': 1, 'submitted': 1}):
            menu.record_application()
        self.assertEqual(seen['default'], radar.today())
        self.assertIn('America/Chicago', seen['prompt'])
        self.assertEqual(radar.read(self.state_path)['applications'][
            self.key_of('Acme', '111')]['date_applied'], radar.today())

    def test_a_backdated_submission_is_accepted(self):
        with mock.patch.object(menu, 'ask', side_effect=['1', '2026-10-05']), \
             mock.patch.object(menu, 'ask_yes_no', side_effect=[True, True]), \
             mock.patch.object(ledger_report, 'build', return_value={
                 'report_path': 'x', 'csv_path': 'y', 'pending': 1, 'submitted': 1}):
            menu.record_application()
        self.assertEqual(radar.read(self.state_path)['applications'][
            self.key_of('Acme', '111')]['date_applied'], '2026-10-05')

    def test_repeat_confirmation_creates_no_second_application(self):
        key = self.key_of('Acme', '111')
        state = radar.read(self.state_path)
        radar.confirm(state, key, '2026-10-05', True)
        radar.write(self.state_path, state)

        with mock.patch.object(menu, 'ask', side_effect=['1']), \
             mock.patch.object(menu, 'ask_yes_no', return_value=True):
            menu.record_application()

        apps = radar.read(self.state_path)['applications']
        self.assertEqual(len(apps), 1)
        self.assertEqual(apps[key]['date_applied'], '2026-10-05',
                         'the original date must not be silently overwritten')

    def test_radar_confirm_itself_refuses_a_second_record(self):
        key = self.key_of('Acme', '111')
        state = radar.read(self.state_path)
        self.assertTrue(radar.confirm(state, key, '2026-10-05', True)['counted'])
        again = radar.confirm(state, key, '2026-10-12', True)
        self.assertFalse(again['counted'])
        self.assertEqual(again['reason'], 'Already recorded')
        self.assertEqual(state['applications'][key]['date_applied'], '2026-10-05')


class ReportRebuild(Fixture):
    def _apply(self, company, job_id, day):
        state = radar.read(self.state_path)
        radar.confirm(state, self.key_of(company, job_id), day, True)
        radar.write(self.state_path, state)

    def test_submitted_jobs_leave_the_pending_shortlist(self):
        self._apply('Acme', '111', '2026-10-05')
        out = ledger_report.build(state_path=self.state_path, reports_dir=self.reports)
        self.assertEqual(out['submitted'], 1)
        self.assertEqual(out['pending'], 1)

        import openpyxl
        wb = openpyxl.load_workbook(out['report_path'])
        pending = [str(c) for row in wb['Pending shortlist'].iter_rows(values_only=True)
                   for c in row if c is not None]
        self.assertFalse(any('111' == c for c in pending), 'submitted role must not be pending')
        history = [str(c) for row in wb['Application history'].iter_rows(values_only=True)
                   for c in row if c is not None]
        self.assertTrue(any('111' == c for c in history))

    def test_history_sheet_carries_the_required_columns(self):
        self._apply('Acme', '111', '2026-10-05')
        import openpyxl
        out = ledger_report.build(state_path=self.state_path, reports_dir=self.reports)
        rows = list(openpyxl.load_workbook(out['report_path'])['Application history']
                    .iter_rows(values_only=True))
        header = next(r for r in rows if r and r[0] == 'Company')
        for col in ('Company', 'Role', 'Requisition ID', 'Official URL', 'Submission date',
                    'Confirmation provenance'):
            self.assertIn(col, header)
        data = rows[rows.index(header) + 1]
        self.assertEqual(data[0], 'Acme')
        self.assertEqual(data[4], '2026-10-05')
        self.assertIn('user confirmed on 2026-10-05', data[5])

    def test_rebuilding_later_does_not_refresh_discovery_dates(self):
        """The whole point: a rebuild is a view, not a re-verification."""
        import openpyxl
        first = ledger_report.build(state_path=self.state_path, reports_dir=self.reports,
                                    now=datetime.datetime(2026, 10, 12, 9, 0, 0))
        later = ledger_report.build(state_path=self.state_path, reports_dir=self.reports,
                                    day='2026-12-25',
                                    now=datetime.datetime(2026, 12, 25, 9, 0, 0))

        def found_dates(path):
            ws = openpyxl.load_workbook(path)['Pending shortlist']
            rows = list(ws.iter_rows(values_only=True))
            hdr = next(i for i, r in enumerate(rows) if r and r[0] == 'Company')
            idx = rows[hdr].index('First found')
            return [r[idx] for r in rows[hdr + 1:] if r and r[0]]

        self.assertEqual(found_dates(first['report_path']), found_dates(later['report_path']))
        self.assertEqual(sorted(found_dates(later['report_path'])), ['2026-10-01', '2026-10-02'])

    def test_previous_reports_are_kept_as_snapshots(self):
        a = ledger_report.build(state_path=self.state_path, reports_dir=self.reports,
                                now=datetime.datetime(2026, 10, 12, 9, 0, 0))
        self._apply('Acme', '111', '2026-10-12')
        b = ledger_report.build(state_path=self.state_path, reports_dir=self.reports,
                                now=datetime.datetime(2026, 10, 12, 10, 0, 0))
        self.assertNotEqual(a['report_path'], b['report_path'])
        self.assertTrue(Path(a['report_path']).exists(), 'the older snapshot must survive')

    def test_a_locked_file_produces_a_new_version_not_a_crash(self):
        stem = 'job_radar_ledger_report_2026-10-12_090000'
        taken = self.reports / (stem + '.xlsx')
        taken.write_bytes(b'pretend Excel holds this')
        out = ledger_report.build(state_path=self.state_path, reports_dir=self.reports,
                                  now=datetime.datetime(2026, 10, 12, 9, 0, 0))
        self.assertNotEqual(Path(out['report_path']), taken)
        self.assertIn('_v2', Path(out['report_path']).name)
        self.assertEqual(taken.read_bytes(), b'pretend Excel holds this',
                         'the existing file must not be touched')

    def test_csv_export_matches_the_canonical_ledger(self):
        import csv
        self._apply('Acme', '111', '2026-10-05')
        out = ledger_report.build(state_path=self.state_path, reports_dir=self.reports)
        rows = list(csv.DictReader(open(out['csv_path'], encoding='utf-8-sig')))
        applied = [r for r in rows if r['Status'] == 'APPLIED']
        self.assertEqual(len(applied), 1)
        self.assertEqual(applied[0]['Job ID'], '111')
        self.assertEqual(applied[0]['Date Applied'], '2026-10-05')

    def test_report_builds_in_a_path_containing_spaces(self):
        out = ledger_report.build(state_path=self.state_path, reports_dir=self.reports)
        self.assertIn(' ', str(self.reports))
        self.assertTrue(Path(out['report_path']).exists())


class LauncherQuotesItsPaths(unittest.TestCase):
    def test_every_use_of_the_project_dir_is_quoted(self):
        text = Path('Start_Job_Radar.cmd').read_text(encoding='utf-8')
        self.assertIn('%~dp0', text)
        for line in text.splitlines():
            if '%PROJECT_DIR%' in line and not line.strip().startswith('REM'):
                if 'set "PROJECT_DIR' in line or '%PROJECT_DIR:~' in line:
                    continue
                self.assertIn('"%PROJECT_DIR%', line,
                              'unquoted path breaks on spaces: %s' % line.strip())

    def test_it_pauses_so_errors_stay_readable(self):
        text = Path('Start_Job_Radar.cmd').read_text(encoding='utf-8')
        self.assertIn('pause', text)
        self.assertIn(':failed', text)
        self.assertIn('cd /d', text)

    def test_it_prefers_a_project_virtualenv(self):
        text = Path('Start_Job_Radar.cmd').read_text(encoding='utf-8')
        self.assertLess(text.index('.venv\\Scripts\\python.exe'), text.index('where py'),
                        'a project venv must win over whatever python is on PATH')


class MenuWiring(unittest.TestCase):
    def test_every_option_is_reachable(self):
        self.assertEqual(sorted(menu.ACTIONS), ['1', '2', '3', '4', '5'])
        for n in '12345':
            self.assertIn('%s.' % n, menu.MENU)
        self.assertIn('0. Exit', menu.MENU)

    def test_only_option_one_spends_money(self):
        for name in ('record_application', 'view_history', 'rebuild_report',
                     'open_latest_report'):
            src = Path('menu.py').read_text(encoding='utf-8')
            self.assertNotIn('agent.py', src.split('def %s' % name)[1].split('\ndef ')[0],
                             '%s must not invoke the paid agent' % name)

    def test_zero_exits_cleanly(self):
        with mock.patch('builtins.input', side_effect=['0']):
            self.assertEqual(menu.main(), 0)

    def test_an_action_error_returns_to_the_menu_instead_of_closing(self):
        with mock.patch.dict(menu.ACTIONS, {'3': mock.Mock(side_effect=RuntimeError('boom'))}), \
             mock.patch('builtins.input', side_effect=['3', '', '0']):
            self.assertEqual(menu.main(), 0)

    def test_unrecognised_input_is_re_prompted(self):
        with mock.patch('builtins.input', side_effect=['zzz', '0']):
            self.assertEqual(menu.main(), 0)


class YesNoIsStrict(unittest.TestCase):
    def test_typos_are_re_asked_not_treated_as_yes(self):
        with mock.patch.object(menu, 'ask', side_effect=['maybe', 'yep', 'y']):
            self.assertTrue(menu.ask_yes_no('ok?'))

    def test_plain_enter_takes_the_stated_default(self):
        with mock.patch.object(menu, 'ask', side_effect=['n']):
            self.assertFalse(menu.ask_yes_no('ok?', default=False))


if __name__ == '__main__':
    unittest.main(verbosity=2)


class PipedInputIsNormalised(unittest.TestCase):
    """A piped "0" arrived with a BOM during launcher testing and was treated
    as an unrecognised option. A human typing never produces one."""

    def test_bom_and_zero_width_are_stripped(self):
        self.assertEqual(menu.clean('﻿0'), '0')
        self.assertEqual(menu.clean('  4 ​'), '4')
        self.assertEqual(menu.clean(None), '')

    def test_a_bom_prefixed_zero_still_exits(self):
        with mock.patch('builtins.input', side_effect=['﻿0']):
            self.assertEqual(menu.main(), 0)
