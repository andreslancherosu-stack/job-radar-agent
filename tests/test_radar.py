"""Tests for the discovery-and-report run: ceilings, honest partials, and the
gates that stop a selection from resting on something the employer never wrote.

No network and no API key. Every HTTP boundary is stubbed, so these prove
MECHANICS -- limits, validation, accounting, file output. They prove nothing
about model judgment; only a real API run does that.
"""
import json
import types
import unittest
from pathlib import Path
from unittest import mock

import agent
import discovery
import radar_tools


def tool_use(name, _id, **inp):
    return types.SimpleNamespace(type='tool_use', id=_id, name=name, input=inp)


def text_block(t):
    return types.SimpleNamespace(type='text', text=t)


class FakeClient:
    """Replays a scripted list of responses, one per API call."""

    def __init__(self, responses, usage=(100, 20)):
        self.responses = list(responses)
        self.calls = 0
        self.usage = usage
        self.messages = self

    def create(self, **kwargs):
        self.calls += 1
        content = self.responses[min(self.calls - 1, len(self.responses) - 1)]
        return types.SimpleNamespace(
            content=content,
            usage=types.SimpleNamespace(input_tokens=self.usage[0], output_tokens=self.usage[1]),
            stop_reason='tool_use' if any(getattr(b, 'type', None) == 'tool_use' for b in content)
                        else 'end_turn')


class IterationCeilingIsEnforced(unittest.TestCase):
    """The professor's hard constraint. A ceiling that is documented but not
    enforced is worth nothing, so assert the ACTUAL number of API calls."""

    def test_loop_never_exceeds_max_steps(self):
        forever = [tool_use('noop', 'x')]
        client = FakeClient([forever])
        result = agent.run_agent('t', client=client, max_steps=25, verbose=False,
                                 tool_definitions=[], tool_registry={'noop': lambda: {'ok': 1}})
        self.assertEqual(result['status'], 'max_steps_reached')
        self.assertEqual(client.calls, 25, 'must stop at exactly the ceiling')
        self.assertEqual(result['metrics']['api_calls'], 25)
        self.assertIn('INCOMPLETE', result['problem'])

    def test_default_radar_ceiling_is_25(self):
        self.assertEqual(agent.RADAR_MAX_STEPS, 25)

    def test_a_ceiling_hit_is_never_reported_as_completed(self):
        client = FakeClient([[tool_use('noop', 'x')]])
        result = agent.run_agent('t', client=client, max_steps=3, verbose=False,
                                 tool_definitions=[], tool_registry={'noop': lambda: {'ok': 1}})
        self.assertNotEqual(result['status'], 'completed')
        self.assertIs(result['required_outcome_met'], None)


class TimeBudgetStopsAndPreservesWork(unittest.TestCase):
    def test_budget_is_checked_before_spending_another_call(self):
        client = FakeClient([[tool_use('noop', 'x')]])
        # Clock starts at 0 so Metrics records started=0, then jumps past the
        # budget and STAYS there. A fixed side_effect list runs out and raises
        # StopIteration from inside the code under test, which would make this
        # test fail for a reason that has nothing to do with the budget.
        ticks = iter([0])

        def clock():
            return next(ticks, 999.0)

        with mock.patch.object(agent.time, 'monotonic', clock):
            result = agent.run_agent('t', client=client, max_steps=25, verbose=False,
                                     time_budget_seconds=10,
                                     tool_definitions=[], tool_registry={'noop': lambda: {'ok': 1}})
        self.assertEqual(result['status'], 'time_budget_exceeded')
        self.assertEqual(client.calls, 0, 'must not spend a call it cannot afford')
        self.assertIn('INCOMPLETE', result['problem'])

    def test_incomplete_run_still_saves_partial_work(self):
        saved = {}

        def on_incomplete(reason):
            saved['reason'] = reason
            return {'report_path': 'x_PARTIAL.xlsx', 'partial': True}

        client = FakeClient([[tool_use('noop', 'x')]])
        result = agent.run_agent('t', client=client, max_steps=2, verbose=False,
                                 on_incomplete=on_incomplete,
                                 tool_definitions=[], tool_registry={'noop': lambda: {'ok': 1}})
        self.assertEqual(saved['reason'], 'max_steps_reached')
        self.assertTrue(result['partial_save']['partial'])

    def test_a_failing_partial_save_does_not_mask_the_real_stop(self):
        def boom(reason):
            raise IOError('disk full')

        client = FakeClient([[tool_use('noop', 'x')]])
        result = agent.run_agent('t', client=client, max_steps=1, verbose=False,
                                 on_incomplete=boom,
                                 tool_definitions=[], tool_registry={'noop': lambda: {'ok': 1}})
        self.assertEqual(result['status'], 'max_steps_reached')
        self.assertIn('disk full', result['partial_save_error'])


class RequestsAndRetriesAreReportedSeparately(unittest.TestCase):
    """api_calls is loop iterations; api_requests is HTTP attempts. Collapsing
    them would hide either wasted spend or the true iteration count."""

    def test_a_retried_step_shows_more_requests_than_calls(self):
        import anthropic

        class FlakyOnce:
            def __init__(self):
                self.n = 0
                self.messages = self

            def create(self, **kw):
                self.n += 1
                if self.n == 1:
                    raise anthropic.APIConnectionError(request=types.SimpleNamespace())
                return types.SimpleNamespace(
                    content=[text_block('done')],
                    usage=types.SimpleNamespace(input_tokens=5, output_tokens=5),
                    stop_reason='end_turn')

        result = agent.run_agent('t', client=FlakyOnce(), max_steps=5, verbose=False,
                                 retry_backoff_seconds=0, tool_definitions=[], tool_registry={})
        m = result['metrics']
        self.assertEqual(m['api_calls'], 1)
        self.assertEqual(m['api_requests'], 2)
        self.assertEqual(m['api_retries'], 1)


class CacheBreakpointsStayUnderTheApiLimit(unittest.TestCase):
    """Measured failure (2026-10-11): the loop set a cache_control breakpoint on
    each turn's newest tool_result but only cleared the CURRENT turn's blocks,
    so old breakpoints accumulated. Step 5 died with 'A maximum of 4 blocks
    with cache_control may be provided. Found 5.' The budget is counted across
    the entire request -- system prompt included."""

    def test_never_more_than_two_breakpoints_across_the_whole_request(self):
        seen = []

        class Recorder:
            def __init__(self):
                self.n = 0
                self.messages = self

            def create(self, **kw):
                blocks = 0
                system = kw.get('system')
                if isinstance(system, list):
                    blocks += sum(1 for b in system if isinstance(b, dict) and b.get('cache_control'))
                for m in kw.get('messages', []):
                    c = m.get('content')
                    if isinstance(c, list):
                        blocks += sum(1 for b in c if isinstance(b, dict) and b.get('cache_control'))
                seen.append(blocks)
                self.n += 1
                content = ([tool_use('noop', 'id%d' % self.n)] if self.n < 7
                           else [text_block('done')])
                return types.SimpleNamespace(
                    content=content,
                    usage=types.SimpleNamespace(input_tokens=5, output_tokens=5,
                                                cache_creation_input_tokens=10,
                                                cache_read_input_tokens=20),
                    stop_reason='tool_use' if self.n < 7 else 'end_turn')

        agent.run_agent('t', client=Recorder(), max_steps=10, verbose=False,
                        tool_definitions=[], tool_registry={'noop': lambda: {'ok': 1}})
        self.assertGreaterEqual(len(seen), 6, 'need enough turns to accumulate breakpoints')
        self.assertLessEqual(max(seen), 4, 'API hard limit is 4; saw %r' % seen)
        self.assertEqual(max(seen), 2, 'system + newest tool_result is all that is needed')

    def test_cache_tokens_are_counted_and_priced_separately(self):
        class C:
            def __init__(self):
                self.messages = self

            def create(self, **kw):
                return types.SimpleNamespace(
                    content=[text_block('done')],
                    usage=types.SimpleNamespace(input_tokens=100, output_tokens=10,
                                                cache_creation_input_tokens=1000,
                                                cache_read_input_tokens=5000),
                    stop_reason='end_turn')

        r = agent.run_agent('t', client=C(), max_steps=2, verbose=False,
                            tool_definitions=[], tool_registry={})
        m = r['metrics']
        self.assertEqual(m['cache_write_tokens'], 1000)
        self.assertEqual(m['cache_read_tokens'], 5000)
        self.assertIn('cache read', m['pricing_note'])

    def test_cache_reads_are_cheaper_than_fresh_input(self):
        fresh, _ = agent.estimate_cost('claude-opus-5', 10_000, 0)
        cached, _ = agent.estimate_cost('claude-opus-5', 0, 0, cache_read_tokens=10_000)
        self.assertLess(cached, fresh / 5)


class CostReportingIsHonest(unittest.TestCase):
    def test_known_model_is_priced(self):
        usd, note = agent.estimate_cost('claude-opus-5', 1_000_000, 1_000_000)
        self.assertAlmostEqual(usd, 90.0)
        self.assertIn('verify against console billing', note)

    def test_unknown_model_reads_as_unknown_not_zero(self):
        usd, note = agent.estimate_cost('some-future-model', 1000, 1000)
        self.assertIsNone(usd, 'an unknown price must never render as $0.00')
        self.assertIn('no list price', note)

    def test_clipped_observation_announces_itself(self):
        out = agent._fmt_observation({'k': 'x' * 5000})
        self.assertIn('trace display clipped', out)
        self.assertIn('the MODEL received', out)


class SelectionGates(unittest.TestCase):
    def setUp(self):
        radar_tools.reset_run()
        radar_tools.RUN.candidates['acme:1'] = {
            'source': 'acme', 'job_id': '1', 'company': 'Acme', 'title': 'Data Analyst'}
        radar_tools.RUN.details['acme:1'] = {
            'company': 'Acme', 'title': 'Data Analyst', 'job_id': '1',
            'url': 'https://example.invalid/1', 'location': 'Dallas'}
        radar_tools.RUN.full_text['acme:1'] = (
            'We seek an analyst.\nBachelor of Science required.\nSQL and Python preferred.')

    def _ok_history(self):
        return mock.patch.object(radar_tools.core_tools, 'execute_tool',
                                 return_value=({'already_applied': False}, False))

    def test_verbatim_quote_is_required(self):
        with self._ok_history():
            out = radar_tools.select_opportunity(
                'acme:1', 1, 'good match', 'Bachelors degree is needed')
        self.assertIn('not found verbatim', out['error'])

    def test_real_quote_is_accepted(self):
        with self._ok_history():
            out = radar_tools.select_opportunity(
                'acme:1', 1, 'good match', 'Bachelor of Science required.')
        self.assertTrue(out['recorded'])
        self.assertEqual(out['remaining_slots'], radar_tools.MAX_SELECTIONS - 1)

    def test_quote_matching_ignores_whitespace_reflow_only(self):
        with self._ok_history():
            out = radar_tools.select_opportunity(
                'acme:1', 1, 'ok', 'SQL and    Python preferred.')
        self.assertTrue(out['recorded'], 'reflowed whitespace is the same quote')

    def test_already_applied_requisition_is_blocked(self):
        with mock.patch.object(radar_tools.core_tools, 'execute_tool',
                               return_value=({'already_applied': True,
                                              'application_record': {'date_applied': '2026-10-04'}},
                                             False)):
            out = radar_tools.select_opportunity(
                'acme:1', 1, 'good', 'Bachelor of Science required.')
        self.assertIn('DUPLICATE', out['error'])
        self.assertEqual(len(radar_tools.RUN.selections), 0)

    def test_cannot_select_a_posting_never_retrieved(self):
        out = radar_tools.select_opportunity('ghost:9', 1, 'x', 'some quote here')
        self.assertIn('get_posting_details', out['error'])

    def test_selection_cap_cannot_be_exceeded(self):
        for i in range(radar_tools.MAX_SELECTIONS):
            k = 'acme:%d' % i
            radar_tools.RUN.details[k] = dict(radar_tools.RUN.details['acme:1'], job_id=str(i))
            radar_tools.RUN.full_text[k] = radar_tools.RUN.full_text['acme:1']
            with self._ok_history():
                radar_tools.select_opportunity(k, i, 'ok', 'Bachelor of Science required.')
        radar_tools.RUN.details['acme:over'] = radar_tools.RUN.details['acme:1']
        radar_tools.RUN.full_text['acme:over'] = radar_tools.RUN.full_text['acme:1']
        with self._ok_history():
            out = radar_tools.select_opportunity('acme:over', 9, 'ok', 'Bachelor of Science required.')
        self.assertIn('maximum', out['error'])
        self.assertEqual(len(radar_tools.RUN.selections), radar_tools.MAX_SELECTIONS)


class CoverageIsReportedHonestly(unittest.TestCase):
    def test_a_dead_source_is_reported_not_silently_dropped(self):
        src = {'a': {'key': 'a', 'company': 'A', 'ats': 'greenhouse', 'board_token': 'a'},
               'b': {'key': 'b', 'company': 'B', 'ats': 'greenhouse', 'board_token': 'b'}}

        def fake(s, timeout):
            if s['key'] == 'b':
                raise OSError('timed out')
            return ([{'source': 'a', 'company': 'A', 'title': 'Data Analyst',
                      'job_id': '1', 'location': 'Dallas', 'url': 'u', 'updated': None,
                      'deadline': None}], 1)

        with mock.patch.object(discovery, '_search_greenhouse', side_effect=fake):
            rows, cov = discovery.search_all(['a', 'b'], keywords=['analyst'], sources=src)
        by = {c['source']: c for c in cov}
        self.assertTrue(by['a']['reached'])
        self.assertFalse(by['b']['reached'])
        self.assertIn('timed out', by['b']['error'])
        self.assertEqual(len(rows), 1, 'a dead source must not abort the others')

    def test_unknown_source_key_is_surfaced(self):
        rows, cov = discovery.search_all(['nope'], sources={})
        self.assertFalse(cov[0]['reached'])
        self.assertIn('unknown source key', cov[0]['error'])

    def test_filter_is_mechanical_and_documented_as_retrieval_only(self):
        row = {'title': 'Senior Data Analyst', 'location': 'Irving, TX'}
        self.assertTrue(discovery.matches(row, ['analyst'], ['irving']))
        self.assertFalse(discovery.matches(row, ['nurse'], None))
        self.assertFalse(discovery.matches(row, None, ['boston']))
        self.assertTrue(discovery.matches(row, None, None), 'no filter keeps everything')


class CoverageAccumulatesAcrossSearches(unittest.TestCase):
    """Measured bug (2026-10-11): a second, narrower search overwrote
    RUN.coverage, so the report's 'actual search coverage' sheet listed one
    source and silently under-reported four employers that had been searched.
    Coverage is a claim about the whole run."""

    def setUp(self):
        radar_tools.reset_run()

    def _c(self, source, reached=True, returned=10, kept=3, error=None, total=100):
        return {'source': source, 'company': source.title(), 'ats': 'workday',
                'reached': reached, 'error': error, 'returned_by_source': returned,
                'total_reported_by_source': total, 'kept_after_filter': kept}

    def test_second_narrower_search_does_not_erase_the_first(self):
        radar_tools.RUN.record_coverage([self._c('gartner'), self._c('jll'), self._c('momentum')])
        radar_tools.RUN.record_coverage([self._c('gartner', returned=5, kept=2)])
        got = {c['source']: c for c in radar_tools.RUN.coverage}
        self.assertEqual(set(got), {'gartner', 'jll', 'momentum'})
        self.assertEqual(got['gartner']['returned_by_source'], 15, 'totals add up')
        self.assertEqual(got['gartner']['kept_after_filter'], 5)
        self.assertEqual(got['gartner']['searches'], 2)
        self.assertEqual(got['jll']['searches'], 1)

    def test_a_source_that_later_succeeds_is_no_longer_marked_failed(self):
        radar_tools.RUN.record_coverage([self._c('jll', reached=False, returned=0, kept=0,
                                                 error='HTTP 500', total=None)])
        radar_tools.RUN.record_coverage([self._c('jll')])
        got = radar_tools.RUN.coverage[0]
        self.assertTrue(got['reached'])
        self.assertIsNone(got['error'])

    def test_a_source_that_never_succeeds_keeps_its_error(self):
        radar_tools.RUN.record_coverage([self._c('jll', reached=False, returned=0, kept=0,
                                                 error='HTTP 500', total=None)])
        radar_tools.RUN.record_coverage([self._c('jll', reached=False, returned=0, kept=0,
                                                 error='timeout', total=None)])
        got = radar_tools.RUN.coverage[0]
        self.assertFalse(got['reached'])
        self.assertIsNotNone(got['error'])


class ExcerptsDeclareWhatTheyOmit(unittest.TestCase):
    def test_sponsorship_sentences_are_never_abbreviated(self):
        text = ('About us.\n'
                'Applicants must be authorized to work in the US without sponsorship now or in '
                'the future.\n' + 'Filler line about culture.\n' * 50)
        out = discovery.extract_requirements(text, excerpt_chars=50)
        self.assertTrue(any('without sponsorship now or in the future' in s
                            for s in out['sponsorship_sentences']))

    def test_omission_is_explicit_and_recoverable(self):
        out = discovery.extract_requirements('Bachelor degree required. ' * 400)
        self.assertIn('are NOT shown', out['omission_note'])
        self.assertIn('full_text=true', out['omission_note'])
        self.assertGreater(out['full_text_chars'], out['shown_chars'])


class ReportIsWrittenAndLabelled(unittest.TestCase):
    def setUp(self):
        radar_tools.reset_run()
        radar_tools.RUN.coverage = [
            {'source': 'gartner', 'company': 'Gartner', 'ats': 'workday', 'reached': True,
             'error': None, 'returned_by_source': 60, 'total_reported_by_source': 178,
             'kept_after_filter': 31},
            {'source': 'dead', 'company': 'Dead', 'ats': 'workday', 'reached': False,
             'error': 'HTTP 500', 'returned_by_source': 0, 'total_reported_by_source': None,
             'kept_after_filter': 0}]

    def _export(self, tmp, **kw):
        with mock.patch.object(radar_tools, 'REPORTS_DIR', Path(tmp)):
            return radar_tools.export_report(**kw)

    def test_partial_run_is_labelled_in_the_filename_and_inside_the_file(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            out = self._export(tmp, partial=True, partial_reason='time_budget_exceeded')
            self.assertIn('_PARTIAL', out['report_path'])
            self.assertTrue(out['partial'])
            wb = openpyxl.load_workbook(out['report_path'])
            rows = list(wb['Selected'].iter_rows(values_only=True))
            status = [r[1] for r in rows if r and r[0] == 'Run status'][0]
            self.assertIn('PARTIAL', status)
            self.assertIn('time_budget_exceeded', status)

    def test_coverage_sheet_records_failures_not_just_successes(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            out = self._export(tmp)
            wb = openpyxl.load_workbook(out['report_path'])
            cells = [str(c) for row in wb['Search coverage'].iter_rows(values_only=True)
                     for c in row if c is not None]
            self.assertIn('HTTP 500', cells)
            self.assertIn('NO', cells)

    def test_report_is_written_even_with_zero_selections(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            out = self._export(tmp)
            self.assertTrue(Path(out['report_path']).exists())
            self.assertEqual(out['selections'], 0)


class TabOpeningIsBoundedAndIdempotent(unittest.TestCase):
    """A measured session left ~15 tabs open. Diagnosis from the real logs:
    each run opened exactly its own shortlist (3+4+0+4+4) -- they accumulated
    ACROSS runs, not within one. Per-run behaviour was already correct, but
    nothing ENFORCED it: a second call, a finalisation retry, or a same-day
    rerun would all have reopened everything."""

    def setUp(self):
        import tempfile
        radar_tools.reset_run()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        p = mock.patch.object(radar_tools, 'REPORTS_DIR', self.dir)
        p.start()
        self.addCleanup(p.stop)
        # a workbook must exist for opening to be allowed at all
        self.report = self.dir / 'r.xlsx'
        self.report.write_bytes(b'x')
        radar_tools.RUN.report_path = str(self.report)
        for i in range(5):
            radar_tools.RUN.selections.append(
                {'key': 'acme:%d' % i, 'company': 'Acme', 'title': 'T%d' % i,
                 'job_id': str(i), 'url': 'https://example.invalid/%d' % i,
                 'priority': i, 'unresolved_questions': []})

    def _open(self, **kw):
        with mock.patch.object(radar_tools.webbrowser, 'open_new_tab',
                               return_value=True) as wb:
            out = radar_tools.open_application_urls(**kw)
        return out, wb

    def test_never_opens_more_than_the_cap(self):
        out, wb = self._open()
        self.assertEqual(out['requested_count'], radar_tools.MAX_SELECTIONS)
        self.assertEqual(wb.call_count, radar_tools.MAX_SELECTIONS,
                         'five selections must not become five tabs')

    def test_second_call_in_the_same_run_opens_nothing(self):
        self._open()
        out2, wb2 = self._open()
        self.assertEqual(out2['requested_count'], 0)
        self.assertEqual(wb2.call_count, 0, 'a finalisation retry must not reopen')
        self.assertTrue(out2['skipped_already_requested_this_run'] or
                        out2['skipped_already_requested_today'])

    def test_same_day_rerun_does_not_reopen(self):
        self._open()
        radar_tools.reset_run()                     # simulate a brand-new run today
        radar_tools.RUN.report_path = str(self.report)
        radar_tools.RUN.selections.append(
            {'key': 'acme:0', 'company': 'Acme', 'title': 'T0', 'job_id': '0',
             'url': 'https://example.invalid/0', 'priority': 1, 'unresolved_questions': []})
        out, wb = self._open()
        self.assertEqual(wb.call_count, 0, 'already open in the browser from earlier today')
        self.assertEqual(len(out['skipped_already_requested_today']), 1)

    def test_force_reopens_only_when_explicitly_asked(self):
        self._open()
        radar_tools.reset_run()
        radar_tools.RUN.report_path = str(self.report)
        radar_tools.RUN.selections.append(
            {'key': 'acme:0', 'company': 'Acme', 'title': 'T0', 'job_id': '0',
             'url': 'https://example.invalid/0', 'priority': 1, 'unresolved_questions': []})
        out, wb = self._open(force=True)
        self.assertEqual(wb.call_count, 1)
        self.assertEqual(out['requested_count'], 1)

    def test_duplicate_urls_in_the_shortlist_open_once(self):
        radar_tools.RUN.selections = [
            {'key': 'a:1', 'company': 'A', 'title': 'T', 'job_id': '1',
             'url': 'https://example.invalid/same', 'priority': 1, 'unresolved_questions': []},
            {'key': 'a:2', 'company': 'A', 'title': 'T', 'job_id': '2',
             'url': 'https://example.invalid/same', 'priority': 2, 'unresolved_questions': []}]
        out, wb = self._open()
        self.assertEqual(wb.call_count, 1)
        self.assertEqual(out['requested_count'], 1)

    def test_refuses_before_the_report_exists(self):
        radar_tools.RUN.report_path = None
        out, wb = self._open()
        self.assertIn('has not been exported', out['error'])
        self.assertEqual(wb.call_count, 0)

    def test_refuses_when_the_report_path_is_stale(self):
        self.report.unlink()
        out, wb = self._open()
        self.assertIn('has not been exported', out['error'])
        self.assertEqual(wb.call_count, 0)

    def test_os_handoff_is_not_reported_as_a_loaded_page(self):
        out, _ = self._open()
        self.assertIn('os_handoff_accepted', out['requested'][0])
        self.assertIn('NOT evidence that the page', out['note'])
        self.assertNotIn('opened_count', out, 'do not claim pages were opened')

    def test_fewer_than_four_qualifying_opens_fewer(self):
        radar_tools.RUN.selections = radar_tools.RUN.selections[:2]
        out, wb = self._open()
        self.assertEqual(wb.call_count, 2)

    def test_a_browser_failure_is_reported_not_swallowed(self):
        with mock.patch.object(radar_tools.webbrowser, 'open_new_tab',
                               side_effect=OSError('no browser')):
            out = radar_tools.open_application_urls()
        self.assertEqual(out['requested_count'], 0)
        self.assertEqual(len(out['failed']), radar_tools.MAX_SELECTIONS)


class ExportAndOpenSetMustAgree(unittest.TestCase):
    def setUp(self):
        import tempfile
        radar_tools.reset_run()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        p = mock.patch.object(radar_tools, 'REPORTS_DIR', Path(self.tmp.name))
        p.start()
        self.addCleanup(p.stop)
        radar_tools.RUN.coverage = [
            {'source': 's', 'company': 'S', 'ats': 'workday', 'reached': True, 'error': None,
             'returned_by_source': 10, 'total_reported_by_source': 10, 'kept_after_filter': 2}]
        for i in range(3):
            radar_tools.RUN.selections.append(
                {'key': 'a:%d' % i, 'company': 'A', 'title': 'T%d' % i, 'job_id': str(i),
                 'url': 'https://example.invalid/%d' % i, 'priority': i,
                 'why_suitable': 'w', 'evidence_quote': 'q', 'unresolved_questions': ['u'],
                 'next_action': 'n', 'posted': None, 'deadline_field': None,
                 'still_listed': True})

    def test_exact_agreement_after_a_normal_finalisation(self):
        radar_tools.export_report()
        with mock.patch.object(radar_tools.webbrowser, 'open_new_tab', return_value=True):
            radar_tools.open_application_urls()
        v = radar_tools.verify_export_matches_open_set()
        self.assertTrue(v['ok'], v.get('problems'))
        self.assertTrue(v['exact_agreement'])
        self.assertEqual(v['workbook_rows'], 3)

    def test_mismatch_is_detected_not_glossed_over(self):
        radar_tools.export_report()
        radar_tools.RUN.opened_urls = [{'key': 'ghost:9', 'company': 'G',
                                        'url': 'https://example.invalid/not-in-report'}]
        v = radar_tools.verify_export_matches_open_set()
        self.assertFalse(v['ok'])
        self.assertTrue(any('not present in the workbook' in p for p in v['problems']))


class AvailabilityConflictsAreHeldBack(unittest.TestCase):
    """Gartner req 112849 advertised an October 9 deadline, then kept being
    listed with that sentence removed and startDate moved to October 10. That
    is consistent with a repost but does not PROVE one, so it must not be
    presented as application-ready -- and no tab may open for it."""

    def setUp(self):
        import tempfile
        radar_tools.reset_run()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        p = mock.patch.object(radar_tools, 'REPORTS_DIR', Path(self.tmp.name))
        p.start()
        self.addCleanup(p.stop)
        radar_tools.RUN.coverage = [
            {'source': 'gartner', 'company': 'Gartner', 'ats': 'workday', 'reached': True,
             'error': None, 'returned_by_source': 60, 'total_reported_by_source': 275,
             'kept_after_filter': 30, 'searches': 1}]
        base = {'company': 'X', 'title': 'T', 'url': 'https://example.invalid/',
                'priority': 1, 'why_suitable': 'w', 'evidence_quote': 'q',
                'unresolved_questions': [], 'next_action': 'n', 'posted': None,
                'deadline_field': None, 'still_listed': True, 'verified_at': 'now'}
        radar_tools.RUN.selections = [
            dict(base, key='ok:1', job_id='1', url='https://example.invalid/ok',
                 availability_verified=True),
            dict(base, key='gartner:112849', job_id='112849',
                 url='https://example.invalid/gartner', availability_verified=False,
                 availability_note='deadline passed; repost not established'),
        ]

    def test_the_real_conflict_file_flags_gartner_112849(self):
        c = radar_tools.availability_conflict('gartner:112849')
        self.assertIsNotNone(c, 'the shipped conflicts file must flag this requisition')
        self.assertEqual(c['classification'], 'availability_unverified')
        self.assertEqual(c['excluded_from'], 'application_ready_shortlist')

    def test_unverified_selection_is_not_application_ready(self):
        self.assertEqual([s['key'] for s in radar_tools.application_ready()], ['ok:1'])
        self.assertEqual([s['key'] for s in radar_tools.needs_verification()],
                         ['gartner:112849'])

    def test_no_tab_opens_for_an_unverified_requisition(self):
        out = radar_tools.export_report()
        radar_tools.RUN.report_path = out['report_path']
        with mock.patch.object(radar_tools.webbrowser, 'open_new_tab',
                               return_value=True) as wb:
            res = radar_tools.open_application_urls()
        urls = [c.args[0] for c in wb.call_args_list]
        self.assertEqual(urls, ['https://example.invalid/ok'])
        self.assertNotIn('https://example.invalid/gartner', urls)
        self.assertEqual(res['requested_count'], 1)

    def test_it_is_kept_in_the_workbook_not_deleted(self):
        import openpyxl
        out = radar_tools.export_report()
        wb = openpyxl.load_workbook(out['report_path'])
        self.assertIn('Needs Verification', wb.sheetnames)
        cells = [str(c) for row in wb['Needs Verification'].iter_rows(values_only=True)
                 for c in row if c is not None]
        self.assertTrue(any('112849' in c for c in cells), 'must be kept, not dropped')
        selected = [str(c) for row in wb['Selected'].iter_rows(values_only=True)
                    for c in row if c is not None]
        self.assertFalse(any('112849' in c for c in selected),
                         'must NOT appear on the application-ready sheet')
        self.assertEqual(out['application_ready'], 1)
        self.assertEqual(out['needs_verification'], 1)

    def test_record_survives_the_model_not_selecting_it(self):
        """The requirement is 'keep its record in Needs Verification'. If the
        sheet only listed held-back SELECTIONS, the record would disappear on
        any run where the model simply did not pick it -- which is exactly what
        happened on the first rehearsal."""
        import openpyxl
        radar_tools.RUN.selections = [s for s in radar_tools.RUN.selections
                                      if s['key'] != 'gartner:112849']
        radar_tools.RUN.candidates['gartner:112849'] = {
            'source': 'gartner', 'job_id': '112849', 'company': 'Gartner',
            'title': 'Data Analyst - 2026/2027 graduates',
            'url': 'https://example.invalid/g112849'}
        out = radar_tools.export_report()
        self.assertEqual(out['needs_verification'], 0, 'it was not selected this run')
        cells = [str(c) for row in openpyxl.load_workbook(out['report_path'])['Needs Verification']
                 .iter_rows(values_only=True) for c in row if c is not None]
        self.assertTrue(any('112849' in c for c in cells),
                        'the record must persist even when unselected')

    def test_scanned_and_evaluated_are_reported_as_different_numbers(self):
        radar_tools.RUN.details = {'gartner:1': {}, 'gartner:2': {}}
        out = radar_tools.export_report()
        self.assertEqual(out['scanned'], 60)
        self.assertEqual(out['fully_evaluated'], 2)
        self.assertNotEqual(out['scanned'], out['fully_evaluated'])


class RadarToolsetCannotTouchAForm(unittest.TestCase):
    def test_no_submission_and_no_browser_write_tools(self):
        _, registry = agent.build_toolset('radar')
        self.assertFalse(any('submit' in n.lower() for n in registry))
        for forbidden in ('fill_application_fields', 'upload_document', 'advance_to_next_page'):
            self.assertNotIn(forbidden, registry)

    def test_required_outcome_is_the_report(self):
        self.assertIn('export_report', agent.build_toolset('radar')[1])


if __name__ == '__main__':
    unittest.main(verbosity=2)
